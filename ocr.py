"""
ocr.py
------
Optical Character Recognition module.

Responsibilities:
  1. Detect whether a PDF page is scanned (image-based) or text-based.
  2. Extract text from text-based PDFs using pdfplumber.
  3. Extract text from scanned PDFs / images using EasyOCR or Tesseract OCR.
  4. Extract text from DOCX and plain-text files.
  5. Preserve structure: paragraphs, section headings, numbering.

Design decision: EasyOCR is preferred because it natively supports Tamil and
Hindi without separate language pack installation. Tesseract is the fallback.
"""

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional

# Monkey patch for EasyOCR 1.6.2 compatibility with modern Pillow (>=10.0.0)
import PIL.Image
if not hasattr(PIL.Image, "ANTIALIAS"):
    PIL.Image.ANTIALIAS = getattr(PIL.Image, "LANCZOS", getattr(PIL.Image.Resampling, "LANCZOS", 1))

import config
from utils import setup_logging, Timer

logger = setup_logging(__name__)


# ─────────────────────────────────────────────
# DATA MODEL
# ─────────────────────────────────────────────

@dataclass
class PageResult:
    """
    Holds the OCR / extraction result for a single document page.

    Attributes:
        page_number:  1-based page index.
        text:         Extracted text content.
        is_ocr:       True if OCR was used (page was scanned).
        confidence:   Mean OCR confidence (0-100) if applicable, else None.
    """
    page_number: int
    text:        str
    is_ocr:      bool = False
    confidence:  Optional[float] = None


@dataclass
class DocumentOCRResult:
    """
    Aggregated extraction result for an entire document.

    Attributes:
        file_path:    Absolute path to the source file.
        file_type:    Extension: 'pdf', 'txt', 'docx'.
        pages:        List of PageResult objects.
        used_ocr:     True if any page required OCR.
        raw_text:     Full concatenated text (all pages joined).
    """
    file_path: Path
    file_type: str
    pages:     list[PageResult] = field(default_factory=list)
    used_ocr:  bool = False

    @property
    def raw_text(self) -> str:
        """Return full document text by joining all pages."""
        return "\n\n".join(p.text for p in self.pages if p.text.strip())

    @property
    def total_pages(self) -> int:
        return len(self.pages)


# ─────────────────────────────────────────────
# LAZY IMPORT HELPERS
# ─────────────────────────────────────────────

def _import_pdfplumber():
    """Lazy import pdfplumber to avoid hard dependency at module load."""
    try:
        import pdfplumber
        return pdfplumber
    except ImportError as e:
        raise ImportError(
            "pdfplumber is required for PDF processing.\n"
            "Install with: pip install pdfplumber"
        ) from e


def _import_easyocr():
    """Lazy import EasyOCR."""
    try:
        import easyocr
        return easyocr
    except ImportError as e:
        raise ImportError(
            "easyocr is required for scanned PDF processing.\n"
            "Install with: pip install easyocr"
        ) from e


def _import_pytesseract():
    """Lazy import pytesseract."""
    try:
        import pytesseract
        return pytesseract
    except ImportError as e:
        raise ImportError(
            "pytesseract is required for Tesseract OCR fallback.\n"
            "Install with: pip install pytesseract"
        ) from e


def _import_pdf2image():
    """Lazy import pdf2image for PDF-to-image conversion."""
    try:
        from pdf2image import convert_from_path
        return convert_from_path
    except ImportError as e:
        raise ImportError(
            "pdf2image is required for scanned PDF processing.\n"
            "Install with: pip install pdf2image\n"
            "Also install Poppler: https://github.com/oschwartz10612/poppler-windows/releases/"
        ) from e


def _import_docx():
    """Lazy import python-docx."""
    try:
        import docx
        return docx
    except ImportError as e:
        raise ImportError(
            "python-docx is required for DOCX processing.\n"
            "Install with: pip install python-docx"
        ) from e


# ─────────────────────────────────────────────
# SCANNED PAGE DETECTION
# ─────────────────────────────────────────────

def _is_page_scanned(page) -> bool:
    """
    Determine if a pdfplumber Page object is scanned (image-based).
    Heuristic: if extracted text has fewer than MIN_TEXT_CHARS_PER_PAGE
    printable characters, treat as scanned.

    Args:
        page: pdfplumber.Page object.

    Returns:
        True if the page is likely scanned.
    """
    text = page.extract_text() or ""
    printable = sum(1 for c in text if c.isprintable() and not c.isspace())
    is_scanned = printable < config.MIN_TEXT_CHARS_PER_PAGE
    logger.debug(
        "Page %d — printable chars: %d → %s",
        page.page_number,
        printable,
        "SCANNED" if is_scanned else "TEXT",
    )
    return is_scanned


# ─────────────────────────────────────────────
# EASYOCR ENGINE
# ─────────────────────────────────────────────

class EasyOCREngine:
    """
    Wrapper around EasyOCR for multilingual text extraction.

    EasyOCR restriction: Tamil, Hindi, and English cannot all be loaded in
    one Reader — they belong to different script groups.

    Solution: three separate cached readers, selected by language code:
        'en' → English-only reader
        'ta' → Tamil + English reader
        'hi' → Hindi  + English reader
    """

    _readers: dict = {}   # Cache: lang_key -> reader instance

    @classmethod
    def _get_reader(cls, lang: str = "en"):
        """
        Return (and cache) the appropriate EasyOCR Reader for the language.

        Args:
            lang: ISO code — 'en', 'ta', or 'hi'.

        Returns:
            easyocr.Reader instance.
        """
        easyocr = _import_easyocr()

        if lang == "ta":
            key   = "ta"
            langs = config.EASYOCR_LANG_TAMIL
        elif lang == "hi":
            key   = "hi"
            langs = config.EASYOCR_LANG_HINDI
        else:
            key   = "en"
            langs = config.EASYOCR_LANG_ENGLISH

        if key not in cls._readers:
            logger.info("Initialising EasyOCR reader for: %s", langs)
            easyocr_model_dir = str(config.MODELS_DIR / "easyocr")

            # Auto-detect GPU — use it if available, else fall back to CPU
            try:
                import torch
                _use_gpu = torch.cuda.is_available()
            except ImportError:
                _use_gpu = False
            logger.info("EasyOCR device: %s", "GPU" if _use_gpu else "CPU")

            try:
                cls._readers[key] = easyocr.Reader(
                    langs, gpu=_use_gpu,
                    model_storage_directory=easyocr_model_dir,
                    quantize=False,
                )
                logger.info("EasyOCR reader [%s] ready.", key)
            except Exception as exc:
                # Model file is corrupted or version-mismatched — delete and retry once
                import glob as _glob, os as _os
                deleted = False
                for bad_file in _glob.glob(_os.path.join(easyocr_model_dir, f"{key}*")) + \
                                _glob.glob(_os.path.join(easyocr_model_dir, "tamil*")):
                    try:
                        _os.remove(bad_file)
                        logger.info("Deleted corrupted model file: %s", bad_file)
                        deleted = True
                    except OSError:
                        pass

                if deleted:
                    logger.info("Retrying EasyOCR reader [%s] after cleanup…", key)
                    try:
                        cls._readers[key] = easyocr.Reader(
                            langs, gpu=_use_gpu,
                            model_storage_directory=easyocr_model_dir,
                            quantize=False,
                        )
                        logger.info("EasyOCR reader [%s] ready after retry.", key)
                        return cls._readers[key]
                    except Exception as exc2:
                        exc = exc2

                logger.warning(
                    "EasyOCR reader [%s] failed (%s). "
                    "Falling back to English-only reader.", key, exc,
                )
                if "en" not in cls._readers:
                    logger.info("Loading English-only EasyOCR reader as fallback.")
                    cls._readers["en"] = easyocr.Reader(
                        config.EASYOCR_LANG_ENGLISH, gpu=_use_gpu,
                        model_storage_directory=easyocr_model_dir,
                        quantize=False,
                    )
                cls._readers[key] = cls._readers["en"]
                logger.info("Fallback English reader will be used for [%s].", key)



        return cls._readers[key]

    @classmethod
    def extract_from_image(cls, image, lang: str = "en") -> tuple[str, float]:
        """
        Run EasyOCR on a PIL Image and return (text, mean_confidence).

        Args:
            image: PIL.Image object.
            lang:  ISO language code of the document ('en', 'ta', 'hi').

        Returns:
            Tuple of (extracted_text, confidence_percent).
        """
        import numpy as np
        reader    = cls._get_reader(lang)
        img_array = np.array(image)
        results   = reader.readtext(img_array, detail=1, paragraph=True)

        lines:       list[str]   = []
        confidences: list[float] = []
        for result in results:
            if len(result) == 3:
                _, text, conf = result
            elif len(result) == 2:
                _, text = result
                conf = 1.0
            else:
                continue
            if text.strip():
                lines.append(text.strip())
                confidences.append(conf * 100)

        full_text = "\n".join(lines)
        mean_conf = sum(confidences) / len(confidences) if confidences else 0.0
        return full_text, mean_conf


# ─────────────────────────────────────────────
# TESSERACT ENGINE (FALLBACK)
# ─────────────────────────────────────────────

class TesseractEngine:
    """Wrapper around pytesseract for OCR extraction."""

    @staticmethod
    def extract_from_image(image) -> tuple[str, float]:
        """
        Run Tesseract OCR on a PIL Image.

        Args:
            image: PIL.Image object.

        Returns:
            Tuple of (extracted_text, mean_confidence).
        """
        pytesseract = _import_pytesseract()
        data = pytesseract.image_to_data(
            image,
            lang=config.TESSERACT_LANG,
            output_type=pytesseract.Output.DICT,
        )
        words       = data["text"]
        confs       = data["conf"]
        valid_words = [
            w for w, c in zip(words, confs)
            if isinstance(c, (int, float)) and int(c) >= 0 and w.strip()
        ]
        valid_confs = [
            int(c) for c in confs
            if isinstance(c, (int, float)) and int(c) >= 0
        ]
        text     = " ".join(valid_words)
        mean_conf = sum(valid_confs) / len(valid_confs) if valid_confs else 0.0
        return text, mean_conf


# ─────────────────────────────────────────────
# MAIN OCR PROCESSOR
# ─────────────────────────────────────────────

class OCRProcessor:
    """
    Main entry point for document text extraction.

    Supports PDF (text-based and scanned), TXT, and DOCX.
    Automatically selects EasyOCR or Tesseract based on config.OCR_BACKEND.
    """

    def __init__(self) -> None:
        self._ocr_engine = (
            EasyOCREngine()
            if config.OCR_BACKEND == "easyocr"
            else TesseractEngine()
        )
        logger.info("OCRProcessor initialised with backend: %s", config.OCR_BACKEND)

    # ── Public API ──────────────────────────────────────────────────────────

    def process(self, file_path: Path) -> DocumentOCRResult:
        """
        Process a document file and return its extracted text.

        Args:
            file_path: Validated Path object to the document.

        Returns:
            DocumentOCRResult containing per-page results and combined text.

        Raises:
            ValueError: If the file type is unsupported.
        """
        suffix = file_path.suffix.lower()
        logger.info("Processing file: %s (type: %s)", file_path.name, suffix)

        with Timer(f"Extraction [{file_path.name}]", logger):
            if suffix == ".pdf":
                return self._process_pdf(file_path)
            elif suffix == ".txt":
                return self._process_txt(file_path)
            elif suffix == ".docx":
                return self._process_docx(file_path)
            else:
                raise ValueError(f"Unsupported file type: {suffix}")

    # ── PDF ─────────────────────────────────────────────────────────────────

    def _process_pdf(self, file_path: Path) -> DocumentOCRResult:
        """Extract text from PDF, using OCR for scanned pages."""
        pdfplumber = _import_pdfplumber()
        result = DocumentOCRResult(file_path=file_path, file_type="pdf")

        with pdfplumber.open(file_path) as pdf:
            total = len(pdf.pages)
            logger.info("PDF has %d page(s).", total)

            # First pass: identify scanned pages
            scanned_indices = [
                i for i, pg in enumerate(pdf.pages)
                if _is_page_scanned(pg)
            ]
            needs_ocr = bool(scanned_indices)
            logger.info(
                "Scanned pages detected: %s",
                scanned_indices if scanned_indices else "none",
            )

            # Detect document language from any available text
            # (used to pick the correct EasyOCR reader)
            ocr_lang = _detect_ocr_language(pdf)
            logger.info("OCR language selected: %s", ocr_lang)

            if needs_ocr:
                # Convert entire PDF to images for OCR pages
                convert_from_path = _import_pdf2image()
                logger.info(
                    "Converting PDF to images at %d DPI for OCR…", config.PDF_OCR_DPI
                )
                images = convert_from_path(
                    str(file_path),
                    dpi=config.PDF_OCR_DPI,
                    poppler_path=config.POPPLER_PATH,  # None = use system PATH
                )
            else:
                images = []

            for i, page in enumerate(pdf.pages):
                page_num = i + 1
                if i in scanned_indices:
                    text, conf = self._ocr_image(images[i], lang=ocr_lang)
                    result.pages.append(
                        PageResult(
                            page_number=page_num,
                            text=text,
                            is_ocr=True,
                            confidence=round(conf, 2),
                        )
                    )
                    result.used_ocr = True
                else:
                    raw = page.extract_text(
                        x_tolerance=3, y_tolerance=3
                    ) or ""
                    result.pages.append(
                        PageResult(
                            page_number=page_num,
                            text=raw,
                            is_ocr=False,
                        )
                    )

        logger.info(
            "PDF extraction complete. Pages: %d, OCR used: %s",
            result.total_pages,
            result.used_ocr,
        )
        return result

    # ── TXT ─────────────────────────────────────────────────────────────────

    def _process_txt(self, file_path: Path) -> DocumentOCRResult:
        """Read plain-text file with automatic encoding detection."""
        result = DocumentOCRResult(file_path=file_path, file_type="txt")
        for encoding in ("utf-8", "utf-8-sig", "latin-1", "cp1252"):
            try:
                text = file_path.read_text(encoding=encoding)
                result.pages.append(PageResult(page_number=1, text=text))
                logger.info(
                    "TXT file read with encoding '%s', %d chars.",
                    encoding, len(text),
                )
                return result
            except UnicodeDecodeError:
                continue
        raise UnicodeDecodeError(
            "utf-8", b"", 0, 1,
            f"Could not decode {file_path.name} with any supported encoding."
        )

    # ── DOCX ────────────────────────────────────────────────────────────────

    def _process_docx(self, file_path: Path) -> DocumentOCRResult:
        """
        Extract text from DOCX preserving paragraph structure.
        Each paragraph becomes a line; tables are extracted row-by-row.
        """
        docx = _import_docx()
        result = DocumentOCRResult(file_path=file_path, file_type="docx")
        doc = docx.Document(str(file_path))

        lines: list[str] = []
        for para in doc.paragraphs:
            if para.text.strip():
                lines.append(para.text)

        # Also extract table contents
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                )
                if row_text:
                    lines.append(row_text)

        full_text = "\n".join(lines)
        result.pages.append(PageResult(page_number=1, text=full_text))
        logger.info(
            "DOCX extraction complete. Paragraphs: %d, total chars: %d",
            len(lines), len(full_text),
        )
        return result

    # ── OCR DISPATCH ─────────────────────────────────────────────────────────

    def _ocr_image(self, image, lang: str = "en") -> tuple[str, float]:
        """
        Dispatch an image to the configured OCR engine.

        Args:
            image: PIL.Image object.
            lang:  Detected document language ('en', 'ta', 'hi').

        Returns:
            (text, confidence_percent) tuple.
        """
        if config.OCR_BACKEND == "easyocr":
            return EasyOCREngine.extract_from_image(image, lang=lang)
        else:
            return TesseractEngine.extract_from_image(image)


# ─────────────────────────────────────────────
# LANGUAGE PRE-DETECTION FOR OCR READER SELECTION
# ─────────────────────────────────────────────

def _detect_ocr_language(pdf) -> str:
    """
    Detect the dominant language of a PDF before full OCR.
    Samples text from the first available text-based page.
    Falls back to 'ta' (Tamil) if no text is found — safe default
    because Tamil+English reader handles English text fine too.

    Args:
        pdf: Open pdfplumber PDF object.

    Returns:
        ISO language code: 'en', 'ta', or 'hi'.
    """
    sample_text = ""
    for page in pdf.pages[:3]:   # Check first 3 pages
        t = page.extract_text() or ""
        if len(t.strip()) > 20:
            sample_text = t
            break

    if not sample_text.strip():
        # No extractable text — assume Tamil (covers Tamil+English docs)
        logger.debug("No text sample for language detection — defaulting to 'ta'.")
        return "ta"

    try:
        from langdetect import detect, DetectorFactory
        DetectorFactory.seed = 42
        code = detect(sample_text)
        if code == "ta":
            return "ta"
        elif code == "hi":
            return "hi"
        else:
            return "en"
    except Exception:
        return "ta"  # Safe fallback
